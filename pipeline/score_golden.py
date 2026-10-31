"""Compare your hand labels (golden set v2) with the AI's answers. Plain code, no AI calls.

    python -m pipeline.score_golden            # reads data/golden/golden_set_v2.xlsx
    python -m pipeline.score_golden --csv      # reads golden_set_v2.csv instead

Run it only AFTER you have filled questions 1 to 4 for all 20 posts. Writes to data/golden/ (git-ignored):
  golden_results.md        agreement per question + every disagreement (post, your answer, AI's answer, your note)
  validation_update.json   the same numbers for the Validation page

How each question is compared:
  1. About SIPs?               vs is_relevant (Yes = relevant).
  2. Writer's own experience?  vs is_own_experience from the t4 run on these 20 posts (tags_t4_golden_*.jsonl); if missing, "not scored" for that post.
  3. What kind of story?       vs lapse_type: involuntary = Payment or app failure; voluntary = Own choice to stop or pause;
                               both = a match with either of those; not_applicable = No stop story; unclear = Unclear.
  4. Main reason               a match if your main reason is anywhere in the AI's reasons (the AI does not rank them);
                               "None given" matches only if the AI gave no reason.
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
G = ROOT / "data" / "golden"
Q = ["1. About SIPs?", "2. Writer's own experience?", "3. What kind of story?", "4. Main reason they give"]
POST = "Post (exactly as written)"
NONE = "None given"
LAPSE_TO_Q3 = {"involuntary": {"Payment or app failure"}, "voluntary": {"Own choice to stop or pause"},
               "both": {"Payment or app failure", "Own choice to stop or pause"}, "not_applicable": {"No stop story"},
               "unclear": {"Unclear"}}
T4_OWN: dict[str, str] = {}


def read_sheet(use_csv: bool) -> list[dict]:
    if use_csv:
        with open(G / "golden_set_v2.csv", newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook
    ws = load_workbook(G / "golden_set_v2.xlsx", data_only=True)["Label here"]
    heads = [str(c.value or "").strip() for c in ws[1]]
    rows = [{h: ("" if v is None else str(v).strip()) for h, v in zip(heads, r)}
            for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    for r in rows:  # Excel / Google Sheets may save the number 1 as 1.0
        r["#"] = str(int(float(r["#"])))
    return rows


def compare(h: dict, a: dict) -> dict:
    """{question: (match True/False/None, my answer, AI's answer)}; None = not scored."""
    raw = a["raw"]
    reasons = [r for r in a["reasons"] if r != NONE]
    out = {Q[0]: (h[Q[0]] == a["about_sips"], h[Q[0]], a["about_sips"])}
    own = T4_OWN.get(a["id"])  # from the t4 run on these 20 posts: "true" / "false" / "unclear"
    if own in ("true", "false", "unclear"):
        shown = {"true": "Yes", "false": "No", "unclear": "Unclear"}[own]
        out[Q[1]] = (h[Q[1]] == shown, h[Q[1]], shown)
    else:
        out[Q[1]] = (None, h[Q[1]], "(AI did not tag this)")
    ok3 = LAPSE_TO_Q3.get(raw["lapse_type"], set())
    out[Q[2]] = (h[Q[2]] in ok3, h[Q[2]], a["choice_or_failure"] + (" (counts as either)" if raw["lapse_type"] == "both" else ""))
    mine = h[Q[3]]
    ok4 = (not reasons) if mine == NONE else (mine in reasons)
    out[Q[3]] = (ok4, mine, ", ".join(reasons) or NONE)
    return out


def main() -> None:
    rows = read_sheet("--csv" in sys.argv)
    ai = json.load(open(G / "ai_answers_v2_hidden.json"))["items"]
    for f in G.glob("tags_t4_golden_*.jsonl"):  # own-experience answers from the single t4 run
        for line in open(f):
            r = json.loads(line)
            if r.get("status") == "ok":
                T4_OWN[r["id"]] = str(r["tags"].get("is_own_experience", "")).lower()
    missing = [(r["#"], q) for r in rows for q in Q if not r.get(q, "").strip()]
    if missing:
        sys.exit(f"{len(missing)} answer cells are still empty (first: post {missing[0][0]}, '{missing[0][1]}'). "
                 "Fill all four questions for every post (use Unclear / None given if the post doesn't say), then rerun.")
    totals = {q: [0, 0] for q in Q}  # agreed, scored
    dis = []
    for r in rows:
        for q, (m, mine, theirs) in compare(r, ai[r["#"]]).items():
            if m is None:
                continue
            totals[q][1] += 1
            totals[q][0] += bool(m)
            if not m:
                dis.append({"#": r["#"], "question": q, "source": r["Source"], "post": r[POST],
                            "my_answer": mine, "ai_answer": theirs, "my_note": r.get("My note", "")})
    table = [{"question": q, "agree": a, "scored": n, "of": len(rows), "pct": round(100 * a / n) if n else None}
             for q, (a, n) in totals.items()]
    md = ["# Golden set v2 results", "", "| Question | Agreed | Scored | Of | % |", "|---|---|---|---|---|"]
    for t in table:
        pct = f"{t['pct']}%" if t["pct"] is not None else "not scored"
        md.append(f"| {t['question']} | {t['agree']} | {t['scored']} | {t['of']} | {pct} |")
    md += ["", f"## Disagreements ({len(dis)})", ""]
    for d in dis:
        md += [f"**Post {d['#']} ({d['source']}), {d['question']}**", f"- Post: {d['post'][:400]}",
               f"- My answer: {d['my_answer']}", f"- AI's answer: {d['ai_answer']}", f"- My note: {d['my_note'] or '(none)'}", ""]
    (G / "golden_results.md").write_text("\n".join(md), encoding="utf-8")
    json.dump({"n_posts": len(rows), "per_question": table, "disagreements": dis},
              open(G / "validation_update.json", "w"), indent=1, ensure_ascii=False)
    for t in table:
        print(f"{t['question']}: {t['agree']} of {t['scored']} scored ({t['pct']}%)" if t["scored"] else f"{t['question']}: not scored")
    print(f"{len(dis)} disagreements. Written to {G}/golden_results.md")


if __name__ == "__main__":
    main()
