"""Proofreading workbook: every post and every AI tag in one spreadsheet you can read, filter and correct.

    python -m pipeline.review_export          # writes data/review/review_workbook.xlsx (local only, git-ignored)
    python -m pipeline.import_review          # reads your REVIEW columns back into data/manual/corrections.csv

Tabs: Summary, All tagged posts (with REVIEW columns and dropdowns), Not yet tagged, Dropped, Allowed values.
Labels are plain English (the same words the dashboard uses). The REVIEW columns hold your corrections; the import step turns
them back into codes. Posts that already have a correction in data/manual/corrections.csv are pre-filled so you can see them.
"""
from __future__ import annotations

import csv
import re
import sys
import warnings
from datetime import datetime

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from collect.common import RAW_DIR, ROOT
from pipeline import export as X
from pipeline.clean import PROCESSED_DIR, staff_ids
from pipeline.tag import CODEBOOK_FILE, load_jsonl, results_for, sample_items

warnings.filterwarnings("ignore")
sys.path.insert(0, str(ROOT / "app"))
import logging  # noqa: E402

logging.disable(logging.CRITICAL)
import common as C  # noqa: E402  (the dashboard's plain-English names, so the workbook and the dashboard use the same words)

REVIEW_DIR = ROOT / "data" / "review"
WORKBOOK = REVIEW_DIR / "review_workbook.xlsx"
CORRECTIONS_FILE = X.CORRECTIONS_FILE
CELL_LIMIT = 32000          # Excel cannot hold more than 32,767 characters in a cell

SEVERITY_LABELS = {1: "1 - Mild annoyance or question", 2: "2 - Caused doubt, a pause, a reduction or a complaint",
                   3: "3 - Caused stopping, money lost, or leaving"}

# review label, field name used by the export / corrections file, codebook field (None = free text or special), kind
#   one = pick one value, many = several values separated by ";", text = free text
FIELDS = [
    ("Relevant", "is_relevant", None, "one"),
    ("Journey stage", "stage", "journey_stage", "one"),
    ("SIP action", "action", "sip_action", "one"),
    ("Lapse type", "lapse_type", "lapse_type", "one"),
    ("Stated triggers", "triggers", "stated_triggers", "many"),
    ("Deeper reason", "reason", None, "text"),
    ("Unmet need", "unmet_need", None, "text"),
    ("Emotion", "emotion", "emotion", "one"),
    ("Severity", "severity", "severity", "one"),
    ("First-time investor", "investor_type", "investor_type", "one"),
    ("Segment income stability", "income", "segment:income_stability", "one"),
    ("Segment age band", "age", "segment:age_band", "one"),
    ("Segment source of money", "money", "segment:money_source", "one"),
    ("Segment who influenced", "influenced_by", "segment:influenced_by", "one"),
    ("Evidence quote", "quote_display", None, "text"),
]
BY_LABEL = {f[0]: f for f in FIELDS}
BY_KEY = {f[1]: f for f in FIELDS}

HEAD_FILL = PatternFill("solid", fgColor="1F3B63")
REVIEW_HEAD_FILL = PatternFill("solid", fgColor="C2570C")
REVIEW_FILL = PatternFill("solid", fgColor="FFF4CC")
THIN = Side(style="thin", color="D9D9D9")


# ---------------------------------------------------------------------------
# labels, in both directions
# ---------------------------------------------------------------------------
def codebook() -> dict:
    return yaml.safe_load(CODEBOOK_FILE.read_text(encoding="utf-8"))["fields"]


def options_of(spec: str | None) -> dict[str, str]:
    """{code: meaning} from codebook.yaml for a codebook field (segment fields are written 'segment:<name>')."""
    if spec is None:
        return {}
    cb = codebook()
    if spec.startswith("segment:"):
        opts = cb["segment_signals"][spec.split(":", 1)[1]]
    else:
        opts = cb[spec]["options"]
    return {str(k): str(v) for k, v in (opts.items() if isinstance(opts, dict) else [(o, "") for o in opts])}


def label_of(key: str, code) -> str:
    """Plain-English label of a stored value for one field."""
    if key == "is_relevant":
        return "Yes" if code in (True, "yes", "Yes", "true", "True") else "No"
    if key == "severity":
        return SEVERITY_LABELS.get(int(code), str(code)) if str(code).isdigit() else str(code)
    if key == "investor_type":
        return {"first_time": "First-time investor", "experienced": "Experienced investor", "unclear": "Not stated"}.get(str(code), str(code))
    if key == "action":
        return C.nice_action(code)
    return C.nice(code)


def allowed_labels(key: str) -> list[str]:
    label, _, spec, kind = BY_KEY[key]
    if key == "is_relevant":
        return ["Yes", "No"]
    if key == "severity":
        return [SEVERITY_LABELS[i] for i in (1, 2, 3)]
    if spec is None:
        return []
    return [label_of(key, code) for code in options_of(spec)]


def code_of(key: str, text: str):
    """Reverse of label_of: the stored value for what you picked or typed (a label, or the raw code). None if not allowed."""
    text = str(text).strip()
    if not text:
        return None
    if key == "is_relevant":
        return {"yes": "yes", "y": "yes", "no": "no", "n": "no"}.get(text.lower())
    if key == "severity":
        return int(text[0]) if text[0] in "123" else None
    spec = BY_KEY[key][2]
    opts = options_of(spec)
    by_label = {label_of(key, c).lower(): c for c in opts}
    return by_label.get(text.lower()) or (text.lower() if text.lower() in opts else None)


# ---------------------------------------------------------------------------
# data
# ---------------------------------------------------------------------------
def read_corrections() -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    if CORRECTIONS_FILE.exists():
        with open(CORRECTIONS_FILE, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                out.setdefault((row.get("id") or "").strip(), []).append(row)
    return out


def year_of(r: dict):
    return r.get("year") or (int(r["created_at"][:4]) if r.get("created_at") else None)


def short(text: str | None) -> str:
    text = (text or "").replace("\x00", "")
    return text if len(text) <= CELL_LIMIT else text[:CELL_LIMIT] + " [cut]"


def segment_text(seg: dict) -> str:
    parts = [(n, seg.get(k)) for n, k in (("Income", "income_stability"), ("Age", "age_band"), ("Money", "money_source"), ("Influenced by", "influenced_by"))]
    stated = [f"{n}: {C.nice(v)}" for n, v in parts if v not in (None, "unclear")]
    return "; ".join(stated) if stated else "none stated"


def quote_status(r: dict) -> str:
    if r.get("quote_unreliable"):
        return "No (hidden on the dashboard)"
    return "Yes, exact" if r.get("quote_exact") else "Yes, repaired in code"


def tagged_rows(records: list[dict], clean: dict, corrections: dict) -> list[dict]:
    rows = []
    for r in records:
        if r["status"] != "ok" or r["id"] not in clean:
            continue
        c, t = clean[r["id"]], r["tags"]
        seg = t["segment_signals"]
        row = {
            "id": r["id"], "source": C.SOURCE_NAMES.get(r["source"], r["source"]), "year": year_of(c), "link": c.get("url"),
            "text": short(c["text"]), "context": short(c.get("context")),
            "relevant": "Yes" if t["is_relevant"] else "No", "stage": label_of("stage", t["journey_stage"]),
            "action": label_of("action", t["sip_action"]), "lapse": label_of("lapse_type", t["lapse_type"]),
            "triggers": "; ".join(label_of("triggers", x) for x in t["stated_triggers"]), "reason": t["underlying_reason"],
            "need": t["unmet_need"], "emotion": label_of("emotion", t["emotion"]), "severity": int(t["severity"]),
            "first_time": label_of("investor_type", t["investor_type"]), "segment": segment_text(seg),
            "quote": r.get("quote_display") or "", "verified": quote_status(r), "by": r.get("tagged_by", "gpt-oss-120b"),
            "needs_review": "Yes" if r.get("needs_review") else "", "_sev": int(t["severity"]), "_year": year_of(c) or 0,
            "review_ok": None, "review_field": None, "review_value": None, "review_note": None,
        }
        done = corrections.get(r["id"])
        if done:                       # show what is already corrected
            first = done[0]
            spec = BY_KEY.get(first["field"])
            if spec:
                vals = [x for x in first["corrected_value"].split("|") if x] if spec[3] == "many" else [first["corrected_value"]]
                row.update(review_ok="N", review_field=spec[0], review_value="; ".join(label_of(first["field"], v) for v in vals) if spec[3] != "text" else first["corrected_value"],
                           review_note=("Already corrected: " + first["reason"]) + (f" (+{len(done) - 1} more correction(s) for this post in corrections.csv)" if len(done) > 1 else ""))
        rows.append(row)
    rows.sort(key=lambda x: (x["source"], -x["_sev"], -x["_year"]))
    return rows


# ---------------------------------------------------------------------------
# writing sheets
# ---------------------------------------------------------------------------
def style_header(ws, review_from: int | None = None) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REVIEW_HEAD_FILL if review_from and cell.column >= review_from else HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put_link(cell, url: str | None) -> None:
    if url:
        cell.value = "open"
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")


def fill_wrap(ws, review_from: int | None = None) -> None:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if review_from and cell.column >= review_from:
                cell.fill = REVIEW_FILL


def sheet_all_tagged(wb: Workbook, rows: list[dict]) -> int:
    ws = wb.create_sheet("All tagged posts")
    heads = ["id", "source", "year", "link", "post text", "context (thread question / video title)", "relevant (tagger)", "journey stage", "SIP action",
             "lapse type", "stated triggers", "deeper reason (AI)", "unmet need (AI)", "emotion", "severity", "first-time investor", "segment signals",
             "evidence quote", "quote verified?", "tagged by", "needs review", "REVIEW: correct? (Y/N)", "REVIEW: field to fix", "REVIEW: correct value", "REVIEW: note"]
    keys = ["id", "source", "year", "link", "text", "context", "relevant", "stage", "action", "lapse", "triggers", "reason", "need", "emotion", "severity",
            "first_time", "segment", "quote", "verified", "by", "needs_review", "review_ok", "review_field", "review_value", "review_note"]
    ws.append(heads)
    for r in rows:
        ws.append([r[k] if k != "link" else None for k in keys])
        put_link(ws.cell(row=ws.max_row, column=4), r["link"])
    review_from = keys.index("review_ok") + 1
    style_header(ws, review_from)
    fill_wrap(ws, review_from)
    set_widths(ws, [24, 14, 7, 7, 70, 42, 10, 16, 18, 16, 30, 34, 30, 13, 28, 16, 28, 44, 16, 14, 10, 12, 26, 34, 40])
    last = max(ws.max_row, 2)
    ok = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showDropDown=False)
    ok.promptTitle, ok.prompt = "Is the AI's tagging correct?", "Y = correct (or leave empty). N = something is wrong: fill in the next three columns."
    field = DataValidation(type="list", formula1=f"'Allowed values'!$A$1:${get_column_letter(len(FIELDS))}$1", allow_blank=True)
    # the value list depends on the field picked in the previous column (named ranges on the 'Allowed values' tab)
    fcol = get_column_letter(review_from + 1)
    value = DataValidation(type="list", formula1=f'INDIRECT("F_"&SUBSTITUTE(SUBSTITUTE({fcol}2," ","_"),"-","_"))', allow_blank=True)
    value.errorStyle, value.showErrorMessage = "information", False
    value.promptTitle, value.prompt = "Correct value", "Pick from the list. For 'Stated triggers' write several values separated by ; . Free-text fields: just type."
    value.showInputMessage = True
    for dv, col in ((ok, review_from), (field, review_from + 1), (value, review_from + 2)):
        dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{last}")
        ws.add_data_validation(dv)
    return ws.max_row - 1


def sheet_allowed(wb: Workbook) -> None:
    """One column of allowed values per field (plain English), with a named range per field for the dependent dropdown."""
    ws = wb.create_sheet("Allowed values")
    ws.append([f[0] for f in FIELDS])
    maxlen = 0
    for ci, (label, key, spec, kind) in enumerate(FIELDS, 1):
        vals = allowed_labels(key)
        maxlen = max(maxlen, len(vals))
        for ri, v in enumerate(vals, 2):
            ws.cell(row=ri, column=ci, value=v)
        if vals:
            name = "F_" + re.sub(r"[ \-]", "_", label)
            col = get_column_letter(ci)
            wb.defined_names[name] = DefinedName(name, attr_text=f"'Allowed values'!${col}$2:${col}${len(vals) + 1}")
        ws.column_dimensions[get_column_letter(ci)].width = 34
    style_header(ws)
    ws.auto_filter.ref = None
    note_row = maxlen + 3
    ws.cell(row=note_row, column=1, value="Free-text fields (Deeper reason, Unmet need, Evidence quote) have no list: type the corrected text. "
                                          "'Stated triggers' can hold several values, separated by ; . If your spreadsheet app does not show the value dropdown "
                                          "(Google Sheets and Numbers may not), type the value exactly as written in the column above.")
    ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)
    ws.row_dimensions[note_row].height = 70


def sheet_untagged(wb: Workbook, clean_rows: list[dict], tagged_ids: set, filter_label: dict, in_sample: set) -> int:
    ws = wb.create_sheet("Not yet tagged")
    ws.append(["id", "source", "year", "link", "post text", "context (thread question / video title)", "quick filter says", "in the tagging sample?"])
    rows = [r for r in clean_rows if r["id"] not in tagged_ids and filter_label.get(r["id"]) != "no"]
    rows.sort(key=lambda r: (C.SOURCE_NAMES.get(r["source"], r["source"]), -(year_of(r) or 0)))
    names = {"yes": "Relevant (yes)", "unsure": "Unsure", None: "Not checked yet"}
    for r in rows:
        ws.append([r["id"], C.SOURCE_NAMES.get(r["source"], r["source"]), year_of(r), None, short(r["text"]), short(r.get("context")),
                   names.get(filter_label.get(r["id"]), filter_label.get(r["id"])), "Yes" if r["id"] in in_sample else "No (outside the sample)"])
        put_link(ws.cell(row=ws.max_row, column=4), r.get("url"))
    style_header(ws)
    fill_wrap(ws)
    set_widths(ws, [24, 14, 7, 7, 80, 42, 16, 22])
    return len(rows)


def sheet_dropped(wb: Workbook, clean_rows: list[dict], tagged_ids: set, filter_info: dict, staff: set) -> int:
    ws = wb.create_sheet("Dropped")
    ws.append(["id", "source", "year", "link", "post text", "removed at", "reason"])
    rows = []
    reasons = {"duplicate": "Duplicate or near-duplicate of another post", "short": "Too short (under 6 words, no SIP word)"}
    for r in load_jsonl(PROCESSED_DIR / "dropped.jsonl"):
        rows.append((r, "Cleaning", reasons.get(r["drop_reason"], r["drop_reason"])))
    for r in clean_rows:
        if r["id"] in staff:
            rows.append((r, "Cleaning", "TradingQ&A staff reply (not an investor; left out of every count, kept only as thread context)"))
        elif filter_info.get(r["id"], {}).get("relevant") == "no" and r["id"] not in tagged_ids:
            why = filter_info[r["id"]].get("reason") or ""
            rows.append((r, "Quick filter", "Judged not about SIPs or mutual funds" + (f": {why}" if why else "")))
    rows.sort(key=lambda t: (t[1], C.SOURCE_NAMES.get(t[0]["source"], t[0]["source"]), -(year_of(t[0]) or 0)))
    for r, step, why in rows:
        ws.append([r["id"], C.SOURCE_NAMES.get(r["source"], r["source"]), year_of(r) or (int(r["created_at"][:4]) if r.get("created_at") else None), None,
                   short(r.get("text")), step, why])
        put_link(ws.cell(row=ws.max_row, column=4), r.get("url"))
    style_header(ws)
    fill_wrap(ws)
    set_widths(ws, [24, 14, 7, 7, 80, 14, 60])
    return len(rows)


def sheet_summary(wb: Workbook, over: dict, counts: dict, n_corr: int) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Review workbook: every post and every AI tag, for proofreading"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Built {datetime.now().strftime('%Y-%m-%d %H:%M')} from the data on this computer (local only, never committed). Data as of {over['as_of']}."
    ws.append([])
    ws.append(["Source", "Pulled", "Cleaned", "Quick-filtered", "Tagged", "Relevant"])
    head_row = ws.max_row
    tot = [0] * 5
    for src, v in over["sources"].items():
        vals = [v["raw_pulled"] or 0, v["clean"], v["filter_labelled"], v["tagged"], v["relevant"]]
        ws.append([C.SOURCE_NAMES.get(src, src)] + vals)
        tot = [a + b for a, b in zip(tot, vals)]
    ws.append(["Total"] + tot)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for cell in ws[head_row]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    ws.append([])
    ws.append(["Tab", "Rows", "What it holds"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    for tab, n, what in (("All tagged posts", counts["tagged"], "Every post the AI has tagged (relevant or not), with its tags in plain English and the REVIEW columns."),
                         ("Not yet tagged", counts["untagged"], "Cleaned posts still waiting for a tagger."),
                         ("Dropped", counts["dropped"], "Posts removed by cleaning (duplicates, too short, staff replies) or by the quick filter, with the reason."),
                         ("Allowed values", "", "The allowed values for each field, used by the dropdowns.")):
        ws.append([tab, n, what])
    ws.append([])
    ws.append(["How to use this workbook"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    steps = [
        "1. Open the tab 'All tagged posts'. Each row is one post with what the AI said. Use the filters on the header row (source, severity, stage, ...).",
        "2. Read the post text and context, then the tags. Rows are sorted by source, then severity (highest first), then newest year first.",
        "3. If everything in a row is right, leave the orange REVIEW columns empty (or write Y).",
        "4. If something is wrong, in the same row: set 'REVIEW: correct? (Y/N)' to N, pick the field in 'REVIEW: field to fix', then pick or type the right value in "
        "'REVIEW: correct value', and say why in 'REVIEW: note'. One correction per row; to fix a second field of the same post, copy the row's id into a free row.",
        "5. Fields you can correct: Relevant, Journey stage, SIP action, Lapse type, Stated triggers (several values separated by ;), Deeper reason, Unmet need, Emotion, "
        "Severity, First-time investor, four segment fields, Evidence quote.",
        "6. Save the file, then run:   python -m pipeline.import_review   It copies each N row into data/manual/corrections.csv, with your note as the reason.",
        "7. Run:   python -m pipeline.export   The corrections are applied everywhere in the dashboard data and listed on the Validation page as human corrections.",
        "Notes: 'Relevant = No' removes the post from every count. Changing a not-relevant post to relevant cannot be applied (it has no tags): tell Claude to re-tag it. "
        f"{n_corr} correction(s) already exist; their rows are pre-filled so you can see them. Staff replies and quick-filter rejects are in the Dropped tab, not in the counts.",
        "Privacy: usernames are never stored. The workbook is in data/review/, which is git-ignored: do not share it publicly.",
    ]
    for s in steps:
        ws.append([s])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=6)
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 18 * (1 + len(s) // 105)
    set_widths(ws, [34, 12, 12, 16, 12, 12])
    ws.column_dimensions["C"].width = 14


# ---------------------------------------------------------------------------
def main() -> None:
    records, by_tagger = X.load_tagged()
    posts = X.build_posts(records, by_tagger)
    over = X.overview(posts, records, by_tagger)
    staff = staff_ids()
    clean_rows = [r for r in load_jsonl(PROCESSED_DIR / "clean.jsonl")]
    clean = {r["id"]: r for r in clean_rows}
    corrections = read_corrections()
    filter_info = {i: v for i, v in results_for("v3").items()}
    filter_label = {i: v["relevant"] for i, v in filter_info.items()}
    tagged_ids = {r["id"] for r in records if r["status"] == "ok"}
    in_sample = {r["id"] for r in sample_items()}

    wb = Workbook()
    rows = tagged_rows(records, clean, corrections)
    counts = {}
    sheet_all_tagged(wb, rows)
    sheet_allowed(wb)
    counts["tagged"] = len(rows)
    live = [r for r in clean_rows if r["id"] not in staff]
    counts["untagged"] = sheet_untagged(wb, live, tagged_ids, filter_label, in_sample)
    counts["dropped"] = sheet_dropped(wb, clean_rows, tagged_ids, filter_info, staff)
    sheet_summary(wb, over, counts, sum(len(v) for v in corrections.values()))
    wb.move_sheet("Allowed values", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Allowed values"))
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)
    print(f"Saved {WORKBOOK.relative_to(ROOT)}: {counts['tagged']} tagged posts, {counts['untagged']} not yet tagged, {counts['dropped']} dropped.")


if __name__ == "__main__":
    main()
