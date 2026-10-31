"""Build the blind golden-set labelling sheet (plain code, NO AI calls, no existing data is changed).

    python -m eval.build_golden_sheet            # seed 3020

Writes to data/golden/ (git-ignored):
  golden_set_to_label.xlsx   tabs: Label here / How to decide / About  (+ a hidden 'Lists' tab for the dropdowns)
  golden_set_to_label.csv    the 'Label here' tab as CSV
  ai_answers_hidden.json     the AI's answers for the same 20 posts. Do not open before labelling.
"""
from __future__ import annotations

import csv
import json
import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "golden"
SEED = 3020

SOURCE_NAMES = {"play_store": "Play Store", "tradingqna": "TradingQ&A", "youtube": "YouTube",
                "reddit": "Reddit", "app_store": "App Store"}
QUOTA = {"play_store": 4, "tradingqna": 4, "youtube": 4, "reddit": 4, "app_store": 2}  # + 2 fill = 20

# ---- plain-English labels for the codebook values --------------------------------------------
ABOUT = ["Yes", "No"]
ACTION = {"started": "Started", "continuing": "Continuing", "considering_stop": "Thinking of stopping",
          "paused": "Paused", "reduced": "Reduced", "stopped": "Stopped", "instalment_failed": "Payment failed",
          "none": "Nothing mentioned"}
LAPSE = {"voluntary": "Person's own choice", "involuntary": "Payment or app failure", "both": "Both",
         "not_applicable": "Not a stop story", "unclear": "Unclear"}
REASON = {"mandate_payment_issue": "Autopay/payment problem", "confusing_numbers": "Confusing numbers (NAV/XIRR)",
          "underwhelming_returns": "Returns too small", "expectation_mismatch": "Expected faster growth",
          "cash_flow_strain": "Money is tight", "irregular_income": "Irregular income",
          "goal_change": "Needs money/goal changed", "social_influence": "Influenced by others",
          "trust_concern": "Distrusts platform", "app_experience": "Hard to use app",
          "market_fall": "Market fell", "life_event": "Life event", "fund_switching": "Switching fund",
          "disengaged_forgot": "Forgot or lost interest", "tax_or_lockin": "Tax or lock-in",
          "other": "Something else", "none": "None given"}
YNU = ["Yes", "No", "Unclear"]
INVESTOR = {"first_time": "Yes", "experienced": "No", "unclear": "Unclear"}

HEAD = ["#", "Source", "Year", "Link", "Post (exactly as written, full text)", "Context (thread question / video title)",
        "About SIPs?", "What happened to the SIP", "Choice or failure", "Main reason", "Other reasons",
        "Writer's own experience?", "First-time investor?", "My note"]
ANSWER_COLS = range(7, 15)  # columns G..N (1-based)

ADVICE = re.compile(r"\b(you should|you can|i would (suggest|recommend)|i suggest|i recommend|better (to|if)|go for|"
                    r"don'?t worry|my advice|you need to|just (do|invest|start)|try (to|the)|make sure)\b", re.I)


def paths_excluded() -> set[str]:
    d = ROOT / "data"
    ex: set[str] = set()
    t = json.load(open(d / "processed/retag_test_ids.json"))  # 50 severe + 50 random hand-checked
    ex |= set(t["severe"]) | set(t["random"])
    ex |= set(json.load(open(d / "processed/pilot_ids.json")))        # relevance pilot, 100
    ex |= set(json.load(open(d / "processed/tag_pilot_ids.json")))    # tagging pilot, 50

    def ids(f):
        return {json.loads(l)["id"] for l in open(f)}
    ex |= ids(d / "processed/tags.jsonl") & ids(d / "processed/tags_qwen.jsonl")  # items tagged by both models (30-post check)
    with open(d / "manual/corrections.csv", newline="", encoding="utf-8") as f:
        ex |= {r["id"] for r in csv.DictReader(f)}
    return ex


def load_candidates(excluded: set[str], min_words: int = 8) -> list[dict]:
    clean = {}
    for l in open(ROOT / "data/processed/clean.jsonl"):
        r = json.loads(l)
        clean[r["id"]] = r
    out = []
    for l in open(ROOT / "data/export/posts.jsonl"):
        p = json.loads(l)
        c = clean.get(p["id"])
        if not c or p["id"] in excluded or int(p["year"]) < 2021:
            continue
        text = (c.get("text") or "").strip()
        if not (min_words <= len(text.split()) and len(text) <= 1200):  # long enough to judge, short enough to read
            continue
        if p["action"] == "increased":  # no matching answer in the dropdown
            continue
        if str(p.get("needs_review")) == "True":
            continue
        p = dict(p)
        p["full_text"], p["full_context"], p["link"] = text, (c.get("context") or "").strip(), c.get("url") or p.get("url")
        trig = set(p["triggers"])
        p["_cats"] = {
            "payment": p["action"] == "instalment_failed" or p["lapse_type"] == "involuntary",
            "chose_stop": p["lapse_type"] == "voluntary",
            "returns": bool(trig & {"confusing_numbers", "underwhelming_returns", "expectation_mismatch"}),
            "plain_q": p["action"] == "none" and "?" in text and trig <= {"none", "other"} and p["lapse_type"] == "not_applicable",
            "hinglish": p["language"] == "hinglish",
            "advice": p["source"] in ("tradingqna", "youtube", "reddit") and bool(ADVICE.search(text)) and p["action"] == "none",
            "first_time": p["investor_type"] == "first_time",
        }
        out.append(p)
    return out


NEED = {"payment": 3, "chose_stop": 3, "returns": 3, "plain_q": 2, "hinglish": 3, "advice": 2, "first_time": 2}


def pick(cands: list[dict], rng: random.Random) -> list[dict] | None:
    chosen: list[dict] = []
    used_src = {s: 0 for s in QUOTA}

    def room(p):
        return used_src[p["source"]] < QUOTA[p["source"]]
    for cat in sorted(NEED, key=lambda c: sum(p["_cats"][c] for p in cands)):  # rarest first
        have = sum(p["_cats"][cat] for p in chosen)
        pool = [p for p in cands if p["_cats"][cat] and p not in chosen and room(p)]
        rng.shuffle(pool)
        while have < NEED[cat] and pool:
            p = pool.pop()
            chosen.append(p)
            used_src[p["source"]] += 1
            have = sum(q["_cats"][cat] for q in chosen)
    for src, q in QUOTA.items():  # fill each source quota
        pool = [p for p in cands if p["source"] == src and p not in chosen]
        rng.shuffle(pool)
        while used_src[src] < q and pool:
            chosen.append(pool.pop())
            used_src[src] += 1
    pool = [p for p in cands if p not in chosen]  # fill to 20
    rng.shuffle(pool)
    while len(chosen) < 20 and pool:
        chosen.append(pool.pop())
    ok = (len(chosen) == 20 and all(sum(p["_cats"][c] for p in chosen) >= n for c, n in NEED.items())
          and all(sum(p["source"] == s for p in chosen) >= q for s, q in QUOTA.items()))
    return chosen if ok else None


def ai_answer(p: dict) -> dict:
    trig = [t for t in p["triggers"] if t] or ["none"]
    own = p.get("is_own_experience")
    return {
        "about_sips": "Yes",
        "what_happened": ACTION[p["action"]],
        "choice_or_failure": LAPSE[p["lapse_type"]],
        "reasons": [REASON[t] for t in trig],  # the AI does not rank its reasons
        "own_experience": {True: "Yes", False: "No"}.get(own, None) if own not in (None, "None") else None,
        "first_time": INVESTOR[p["investor_type"]],
        "raw": {"action": p["action"], "lapse_type": p["lapse_type"], "triggers": trig,
                "investor_type": p["investor_type"], "is_own_experience": own},
    }


# ---- the sheet --------------------------------------------------------------------------------
HOW = [  # (column, answer, plain line, short example)
    ("About SIPs?", "Yes", "The post is about mutual funds or SIPs in some way: starting, continuing, stopping, a payment, returns, the app, or being a new investor.", "\"My SIP of 2000 got debited twice.\""),
    ("About SIPs?", "No", "Clearly nothing to do with mutual funds: only stocks, F&O, IPOs, demat or broker charges, spam, or praise with no content.", "\"Options trading here is superfast, best app.\""),
    ("What happened to the SIP", "Started", "The writer started a SIP (or is setting one up).", "\"Started my first SIP yesterday.\""),
    ("What happened to the SIP", "Continuing", "The writer is still investing.", "\"Two years in, still putting 5k every month.\""),
    ("What happened to the SIP", "Thinking of stopping", "Considering pausing or stopping, not done yet.", "\"Should I stop my SIP? It shows -3%.\""),
    ("What happened to the SIP", "Paused", "Used pause or skip.", "\"I paused my SIP for two months.\""),
    ("What happened to the SIP", "Reduced", "Lowered the amount.", "\"Cut my SIP from 10k to 5k.\""),
    ("What happened to the SIP", "Stopped", "Cancelled or stopped it.", "\"I closed my SIP last week.\""),
    ("What happened to the SIP", "Payment failed", "A debit failed, or the SIP stopped without the person deciding.", "\"Autopay failed 3 times and now my SIP is cancelled.\""),
    ("What happened to the SIP", "Nothing mentioned", "No SIP action in the post (a question, a complaint about something else, general advice).", "\"What is NAV?\""),
    ("Choice or failure", "Person's own choice", "The person chose to pause, reduce or stop.", "\"I stopped because returns were low.\""),
    ("Choice or failure", "Payment or app failure", "Payment failed, mandate expired or rejected, bank changed, low balance, or an app problem. The person did not intend to stop.", "\"Mandate expired, SIP stopped by itself.\""),
    ("Choice or failure", "Both", "Elements of both in the same story.", "\"Account was empty, and honestly I had also lost interest.\""),
    ("Choice or failure", "Not a stop story", "No stopping, pausing or risk of it is mentioned.", "\"How do I download my statement?\""),
    ("Choice or failure", "Unclear", "Cannot tell from the post.", "\"My SIP is not showing.\" (could be either)"),
    ("Main reason", "(any reason on the list)", "Pick the ONE reason the writer themselves give as the biggest. Only the writer's own reasons, not what they advise others. If nothing is given, pick None given. The AI does not rank its reasons, so your main reason counts as a match if the AI listed it anywhere.", ""),
    ("Main reason", "Autopay/payment problem", "Autopay, e-mandate, UPI mandate, bank debit, failed instalment.", "\"SIP not deducted this month.\""),
    ("Main reason", "Confusing numbers (NAV/XIRR)", "Doesn't understand NAV, units, XIRR, absolute vs yearly returns, or statements.", "\"Why does XIRR show 12% when I made 3%?\""),
    ("Main reason", "Returns too small", "Returns are small, flat or slower than they hoped.", "\"Two years and only 4% gain.\""),
    ("Main reason", "Expected faster growth", "Expected quick or guaranteed growth; investing \"not working\".", "\"I thought SIP doubles money in a year.\""),
    ("Main reason", "Money is tight", "Expenses, EMI, rent, tight month.", "\"Can't afford 5k this month.\""),
    ("Main reason", "Irregular income", "Freelance, gig, commission or variable salary timing.", "\"Freelancer, some months nothing comes in.\""),
    ("Main reason", "Needs money/goal changed", "Goal done or changed, or needs to withdraw now.", "\"Need the money for a house deposit.\""),
    ("Main reason", "Influenced by others", "A friend, family member, finfluencer, YouTube or Reddit told them to start, stop or switch.", "\"A YouTuber said stop SIPs in small cap.\""),
    ("Main reason", "Distrusts platform", "Doubts about the platform, fund house or safety of money.", "\"Not sure my money is safe here.\""),
    ("Main reason", "Hard to use app", "Hard to find, change or pause a SIP, bugs, app errors.", "\"Can't find where to pause my SIP.\""),
    ("Main reason", "Market fell", "Portfolio red, crash, loss, because the market fell.", "\"Market crashed, everything is red.\""),
    ("Main reason", "Life event", "Job loss, job change, marriage, medical, relocation, studies.", "\"Lost my job in March.\""),
    ("Main reason", "Switching fund", "Stopping to move to a \"better\" fund.", "\"Stopped this one to start a flexi cap SIP.\""),
    ("Main reason", "Forgot or lost interest", "Stopped looking, forgot, lost interest.", "\"Honestly forgot it exists.\""),
    ("Main reason", "Tax or lock-in", "ELSS lock-in, exit load, tax.", "\"Exit load is eating my redemption.\""),
    ("Main reason", "Something else", "A reason that fits none of the above.", ""),
    ("Main reason", "None given", "The writer gives no reason.", "\"Which SIP is best for beginners?\""),
    ("Other reasons", "(same list)", "Any further reasons the writer gives. Pick one from the dropdown, or type several separated by a semicolon (for example: Money is tight; Market fell). Leave empty if there are none. Counted together with the main reason as the full set.", ""),
    ("Writer's own experience?", "Yes", "The writer describes what happened to THEM.", "\"My SIP failed twice.\""),
    ("Writer's own experience?", "No", "A reply, advice or explanation to someone else, or a general statement; not the writer's own story.", "\"You should check if your mandate is active.\""),
    ("Writer's own experience?", "Unclear", "Cannot tell whether it is their own story.", ""),
    ("First-time investor?", "Yes", "Says or clearly implies this is their first SIP or first year of investing.", "\"I am new to investing, started last month.\""),
    ("First-time investor?", "No", "Mentions years of investing, many funds, past market cycles.", "\"Investing since 2016 through two crashes.\""),
    ("First-time investor?", "Unclear", "Not stated. Most posts are this. Guessing is worse than Unclear.", ""),
    ("Any column", "Golden rule", "Label what the post states or plainly implies. If it gives you nothing to go on, choose the Unclear / None option. Do not guess. Use the My note column for anything odd.", ""),
]

ABOUT_TEXT = [
    f"How the 20 were picked: plain code (no AI calls), fixed random seed {SEED}. Only relevant tagged posts from 2021 onward, with a readable length (8 words to 1,200 characters).",
    "Left out: the 100 hand-checked posts (50 severe + 50 random), the relevance pilot (100), the tagging pilot (50), the 30-post two-model check (and any post tagged by both models), and every post in corrections.csv. A few posts the AI itself flagged for review were also left out.",
    "Mix: 4 each from Play Store, TradingQ&A, YouTube and Reddit, 2 from App Store, 2 more to fill. The AI's tags were used ONLY to build the mix: at least 3 payment-failure, 3 chose-to-stop, 3 returns-confusion, 2 plain questions, 3 Hinglish, 2 replies or advice, 2 first-time investors.",
    "Known limits: this is a small, deliberately mixed sample, not a random one, so agreement here is a check on specific kinds of posts, not an overall accuracy figure. 'Replies or advice' were picked with a word rule on the text, because the AI's own-experience tag does not exist for these posts yet.",
]


def build(rows: list[dict]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Label here"
    ws.append(HEAD)
    table = []
    for i, p in enumerate(rows, 1):
        r = [i, SOURCE_NAMES[p["source"]], int(p["year"]), p["link"] or "(no link)", p["full_text"], p["full_context"]] + [""] * 8
        ws.append(r)
        table.append(r)
    hdr_fill, ans_fill, ans_head = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="FFF4C2"), PatternFill("solid", fgColor="F2B600")
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF" if c not in ANSWER_COLS else "000000")
        cell.fill = hdr_fill if c not in ANSWER_COLS else ans_head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [5, 12, 7, 28, 70, 38, 12, 20, 22, 28, 30, 16, 14, 34]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for r in range(2, len(rows) + 2):
        for c in range(1, len(HEAD) + 1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in ANSWER_COLS:
                cell.fill = ans_fill
        ws.row_dimensions[r].height = min(409, max(60, 15 * (len(rows[r - 2]["full_text"]) // 75 + 2)))
    ws.freeze_panes = "E2"
    ws.row_dimensions[1].height = 34

    lists = wb.create_sheet("Lists")
    cols = {"about": ABOUT, "action": list(ACTION.values()), "lapse": list(LAPSE.values()),
            "reason": list(REASON.values()), "ynu": YNU}
    ref = {}
    for ci, (k, vals) in enumerate(cols.items(), 1):
        for ri, v in enumerate(vals, 1):
            lists.cell(ri, ci, v)
        ref[k] = f"Lists!${get_column_letter(ci)}$1:${get_column_letter(ci)}${len(vals)}"
    lists.sheet_state = "hidden"
    last = len(rows) + 1
    for col, key, strict in [("G", "about", True), ("H", "action", True), ("I", "lapse", True), ("J", "reason", True),
                             ("K", "reason", False), ("L", "ynu", True), ("M", "ynu", True)]:
        dv = DataValidation(type="list", formula1=f"={ref[key]}", allow_blank=True, showErrorMessage=strict)
        if not strict:
            dv.showErrorMessage = False  # lets you type several reasons separated by ;
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    hw = wb.create_sheet("How to decide")
    hw.append(["Column", "Answer", "What it means (the rule the AI was given)", "Short example"])
    for row in HOW:
        hw.append(list(row))
    for c in range(1, 5):
        hw.cell(1, c).font, hw.cell(1, c).fill = Font(bold=True, color="FFFFFF"), hdr_fill
    for c, w in enumerate([26, 30, 90, 48], 1):
        hw.column_dimensions[get_column_letter(c)].width = w
    for r in hw.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    hw.freeze_panes = "A2"

    ab = wb.create_sheet("About")
    for t in ABOUT_TEXT:
        ab.append([t])
    ab.column_dimensions["A"].width = 140
    for r in ab.iter_rows():
        r[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT / "golden_set_to_label.xlsx")

    with open(OUT / "golden_set_to_label.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEAD)
        w.writerows(table)

    hidden = {"seed": SEED, "note": "AI answers for the 20 posts. Do not open before labelling.",
              "items": {str(i): {"id": p["id"], **ai_answer(p)} for i, p in enumerate(rows, 1)}}
    json.dump(hidden, open(OUT / "ai_answers_hidden.json", "w"), indent=1, ensure_ascii=False)


def main() -> None:
    excluded = paths_excluded()
    cands = load_candidates(excluded)
    rng = random.Random(SEED)
    for _ in range(5000):
        chosen = pick(cands, rng)
        if chosen:
            break
    else:
        sys.exit("Could not satisfy the mix with the available posts.")
    rng.shuffle(chosen)  # so the order does not reveal the source or the mix
    build(chosen)
    print(f"excluded {len(excluded)} ids; {len(cands)} candidates; 20 picked.")


if __name__ == "__main__":
    main()
